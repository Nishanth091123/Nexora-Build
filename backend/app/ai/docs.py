"""Document Q&A (RAG): read project documents → chunks → embeddings (pgvector) + keyword index → search.

Text is taken from PDF (text layer), Word (.docx), Excel (.xlsx), CSV and TXT files. Scanned PDFs / photos have no
text layer – the assistant says so (they can be read with Drawing Analysis instead).
Nothing is ever deleted: a file is indexed once; later runs only add missing embeddings.
"""
from __future__ import annotations

import io
import logging
import re
import zipfile

from sqlalchemy import text as sql
from sqlalchemy.orm import Session

from .. import llm
from ..models import AIDocChunk, AIIndexedFile, StoredFile
from ..storage import backend_for
from .textsearch import BM25

log = logging.getLogger("nexora.ai.docs")
READABLE = (".pdf", ".docx", ".xlsx", ".csv", ".txt")
CHUNK = 1000
OVERLAP = 150
MAX_FILES_PER_RUN = 8


def ext_of(name: str) -> str:
    return ("." + name.rsplit(".", 1)[-1].lower()) if "." in name else ""


def readable(f: StoredFile) -> bool:
    return ext_of(f.original_name) in READABLE


# ------------------------------------------------------------------ text extraction
def extract(name: str, data: bytes) -> list[tuple[int, str]]:
    """[(page_no, text)] – page 0 when the format has no pages."""
    ext = ext_of(name)
    if ext == ".pdf":
        from pypdf import PdfReader
        r = PdfReader(io.BytesIO(data))
        out = []
        for i, pg in enumerate(r.pages[:300]):
            try:
                t = pg.extract_text() or ""
            except Exception:
                t = ""
            out.append((i + 1, t))
        return out
    if ext == ".docx":
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            xml = z.read("word/document.xml").decode("utf-8", "ignore")
        xml = re.sub(r"</w:p>", "\n", xml)
        xml = re.sub(r"<w:tab/>", "\t", xml)
        txt = re.sub(r"<[^>]+>", "", xml)
        for a, b in (("&amp;", "&"), ("&lt;", "<"), ("&gt;", ">"), ("&quot;", '"'), ("&apos;", "'")):
            txt = txt.replace(a, b)
        return [(0, txt)]
    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets[:20]:
            lines = []
            for row in ws.iter_rows(values_only=True):
                vals = [str(v) for v in row if v not in (None, "")]
                if vals:
                    lines.append(" | ".join(vals))
                if len(lines) > 3000:
                    break
            out.append((0, f"Sheet {ws.title}\n" + "\n".join(lines)))
        return out
    if ext in (".csv", ".txt"):
        for enc in ("utf-8", "utf-16", "cp1252"):
            try:
                return [(0, data.decode(enc))]
            except UnicodeDecodeError:
                continue
        return [(0, data.decode("utf-8", "ignore"))]
    return []


def chunks_of(pages: list[tuple[int, str]]) -> list[tuple[int, str]]:
    out = []
    for page, t in pages:
        t = re.sub(r"[ \t]+", " ", t or "")
        t = re.sub(r"\n{3,}", "\n\n", t).strip()
        if not t:
            continue
        i = 0
        while i < len(t):
            end = min(len(t), i + CHUNK)
            if end < len(t):  # break at a sentence / line end when possible
                cut = max(t.rfind("\n", i + CHUNK // 2, end), t.rfind(". ", i + CHUNK // 2, end))
                if cut > i:
                    end = cut + 1
            piece = t[i:end].strip()
            if len(piece) > 30:
                out.append((page, piece))
            if end >= len(t):
                break
            i = max(end - OVERLAP, i + 1)
    return out


# ------------------------------------------------------------------ indexing
def _vec(v: list[float]) -> str:
    return "[" + ",".join(f"{x:.6f}" for x in v) + "]"


def _is_pg(db: Session) -> bool:
    return db.get_bind().dialect.name == "postgresql"


def _store_embeddings(db: Session, rows: list[AIDocChunk]) -> bool:
    if not rows or not _is_pg(db):
        return False
    vecs = llm.embed([r.text for r in rows])
    if not vecs:
        return False
    for r, v in zip(rows, vecs):
        db.execute(sql("UPDATE ai_doc_chunks SET embedding = CAST(:v AS vector) WHERE id = :id"), {"v": _vec(v), "id": r.id})
    return True


def index_file(db: Session, f: StoredFile) -> AIIndexedFile:
    rec = db.query(AIIndexedFile).filter_by(file_id=f.id).first()
    if rec:
        return rec
    rec = AIIndexedFile(file_id=f.id, project_id=f.project_id)
    try:
        data = backend_for(f.storage_backend).read(f.storage_key)
        pages = extract(f.original_name, data)
        rec.pages = len(pages)
        parts = chunks_of(pages)
        if not parts:
            rec.status = "empty"
            rec.error = ("No text found – this looks like a scanned drawing or photo. Use Drawing Analysis for plans."
                         if ext_of(f.original_name) == ".pdf" else "No text found in this file.")
        rows = [AIDocChunk(project_id=f.project_id, file_id=f.id, chunk_no=i, page=pg, text=t)
                for i, (pg, t) in enumerate(parts[:1500])]
        db.add_all(rows)
        db.flush()
        rec.chunks = len(rows)
    except Exception as e:
        log.warning("could not read %s: %s", f.original_name, e)
        db.rollback()
        rec = AIIndexedFile(file_id=f.id, project_id=f.project_id, status="failed",
                            error=f"Could not read this file: {str(e)[:200]}")
        rows = []
    db.add(rec)
    db.commit()
    if rows:  # keyword search already works; embeddings are a bonus
        try:
            if _store_embeddings(db, rows):
                rec.embedded = True
            db.commit()
        except Exception as e:
            db.rollback()
            log.info("embedding failed for file %s: %s", f.id, e)
    return rec


def ensure_indexed(db: Session, project_id: int) -> dict:
    """Index new readable files of the project (a few per call) and add missing embeddings."""
    files = (db.query(StoredFile).filter(StoredFile.project_id == project_id, StoredFile.is_hidden.is_(False))
             .order_by(StoredFile.id).all())
    done = {r.file_id: r for r in db.query(AIIndexedFile).filter(AIIndexedFile.project_id == project_id).all()}
    new = [f for f in files if readable(f) and f.id not in done][:MAX_FILES_PER_RUN]
    for f in new:
        done[f.id] = index_file(db, f)
    # embeddings added later (e.g. Gemini key added after upload)
    if llm.keys()["gemini"] and _is_pg(db):
        for rec in [r for r in done.values() if r.status == "indexed" and not r.embedded and r.chunks][:3]:
            rows = db.query(AIDocChunk).filter(AIDocChunk.file_id == rec.file_id).order_by(AIDocChunk.chunk_no).all()
            try:
                if _store_embeddings(db, rows):
                    rec.embedded = True
                    db.commit()
            except Exception as e:
                db.rollback()
                log.info("late embedding failed: %s", e)
                break
    return {"files": len([f for f in files if readable(f)]), "indexed": sum(1 for r in done.values() if r.status == "indexed")}


def status(db: Session, project_id: int) -> list[dict]:
    files = (db.query(StoredFile).filter(StoredFile.project_id == project_id, StoredFile.is_hidden.is_(False))
             .order_by(StoredFile.created_at.desc()).all())
    recs = {r.file_id: r for r in db.query(AIIndexedFile).filter(AIIndexedFile.project_id == project_id).all()}
    out = []
    for f in files:
        r = recs.get(f.id)
        out.append({"file_id": f.id, "name": f.original_name, "category": f.category, "size_bytes": f.size_bytes,
                    "created_at": f.created_at.isoformat(), "readable": readable(f),
                    "status": r.status if r else ("waiting" if readable(f) else "not_text"),
                    "chunks": r.chunks if r else 0, "pages": r.pages if r else 0,
                    "embedded": bool(r and r.embedded), "error": r.error if r else ""})
    return out


# ------------------------------------------------------------------ search
def search(db: Session, project_id: int, query: str, k: int = 6, file_ids: list[int] | None = None) -> list[dict]:
    hidden = {f.id for f in db.query(StoredFile.id).filter(StoredFile.project_id == project_id,
                                                           StoredFile.is_hidden.is_(True)).all()}
    q = db.query(AIDocChunk).filter(AIDocChunk.project_id == project_id)
    if file_ids:
        q = q.filter(AIDocChunk.file_id.in_(file_ids))
    rows = [r for r in q.all() if r.file_id not in hidden]
    if not rows:
        return []
    ranks: dict[int, float] = {}
    bm = BM25([r.text for r in rows])
    for pos, (i, s) in enumerate(bm.top(query, 12, 0.3)):
        ranks[rows[i].id] = ranks.get(rows[i].id, 0) + 1 / (60 + pos)
    if _is_pg(db) and llm.keys()["gemini"]:
        try:
            qv = llm.embed([query], task="RETRIEVAL_QUERY")
            if qv:
                params = {"q": _vec(qv[0]), "p": project_id}
                extra = ""
                if file_ids:
                    extra = " AND file_id = ANY(:f)"
                    params["f"] = list(file_ids)
                res = db.execute(sql("SELECT id, 1 - (embedding <=> CAST(:q AS vector)) AS sim FROM ai_doc_chunks "
                                     "WHERE project_id = :p AND embedding IS NOT NULL" + extra +
                                     " ORDER BY embedding <=> CAST(:q AS vector) LIMIT 12"), params).all()
                for pos, (cid, sim) in enumerate(res):
                    if sim is not None and sim > 0.45 and cid in {r.id for r in rows}:
                        ranks[cid] = ranks.get(cid, 0) + 1 / (60 + pos)
        except Exception as e:
            db.rollback()
            log.info("vector search failed: %s", e)
    by_id = {r.id: r for r in rows}
    names = {f.id: f.original_name for f in db.query(StoredFile).filter(StoredFile.id.in_({r.file_id for r in rows})).all()}
    best = sorted(ranks.items(), key=lambda kv: -kv[1])[:k]
    return [{"chunk_id": cid, "file_id": by_id[cid].file_id, "name": names.get(by_id[cid].file_id, "document"),
             "page": by_id[cid].page, "text": by_id[cid].text, "score": round(s, 4)} for cid, s in best]
