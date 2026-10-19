"""Excel (openpyxl) and PDF (reportlab) exports for a saved estimate."""
from __future__ import annotations

import io
from datetime import datetime

NAVY = "0F2350"
BLUE = "1D5AE0"


def inr(v: float) -> str:
    """Indian digit grouping: 12,34,567"""
    v = round(float(v or 0))
    s = str(abs(int(v)))
    if len(s) > 3:
        head, tail = s[:-3], s[-3:]
        parts = []
        while len(head) > 2:
            parts.insert(0, head[-2:])
            head = head[:-2]
        if head:
            parts.insert(0, head)
        s = ",".join(parts) + "," + tail
    return ("-" if v < 0 else "") + s


def header_lines(project: dict, est: dict) -> list[tuple[str, str]]:
    r = est["result"]
    t = r["totals"]
    return [
        ("Project", f"{project.get('name', '')} ({project.get('code', '')})"),
        ("Location", ", ".join(x for x in [project.get("location"), project.get("district"), project.get("state")] if x)),
        ("Estimate", f"{est.get('title') or 'Estimate'} – version {est.get('version')}"),
        ("Rates", f"{est.get('rate_state', 'Tamil Nadu')} · {r['spec_level'].title()} specification"),
        ("Structure", f"{'RCC framed' if r['structure'] == 'rcc_framed' else 'Load bearing'} · masonry: {r['masonry']}"),
        ("Built-up area", f"{t['builtup_sqft']:,.0f} sq ft ({t['builtup_m2']:,.1f} m²)"),
        ("Total cost", f"Rs. {inr(t['grand_total'])}"),
        ("Cost per sq ft", f"Rs. {inr(t['cost_per_sqft'])}"),
        ("Duration", f"{t['duration_months']} months (approx.)"),
        ("Generated", datetime.now().strftime("%d %b %Y %H:%M")),
    ]


# ------------------------------------------------------------------ Excel
def to_excel(project: dict, est: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    r = est["result"]
    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor=NAVY)
    thin = Side(style="thin", color="D9DEE8")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    money = '#,##,##0'

    def table(ws, start_row, headers, rows, money_cols=(), widths=None):
        for ci, h in enumerate(headers, 1):
            c = ws.cell(row=start_row, column=ci, value=h)
            c.font, c.fill, c.border = head_font, head_fill, border
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for ri, row in enumerate(rows, start_row + 1):
            for ci, v in enumerate(row, 1):
                c = ws.cell(row=ri, column=ci, value=v)
                c.border = border
                if ci in money_cols:
                    c.number_format = money
        if widths:
            for i, w in enumerate(widths, 1):
                ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
        return start_row + len(rows) + 1

    # Summary
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Nexora Build – Construction Estimate"
    ws["A1"].font = Font(bold=True, size=15, color=NAVY)
    row = 3
    for k, v in header_lines(project, est):
        ws.cell(row=row, column=1, value=k).font = Font(bold=True)
        ws.cell(row=row, column=2, value=v)
        row += 1
    row += 1
    t = r["totals"]
    rows = [["Materials", t["materials"]], ["Labour", t["labour"]], ["Equipment", t["equipment"]]]
    rows += [[f"{e['name']} ({e['basis']})", e["amount"]] for e in r["extras"]]
    rows += [["GRAND TOTAL", t["grand_total"]]]
    end = table(ws, row, ["Cost head", "Amount (Rs.)"], rows, money_cols=(2,), widths=[48, 22])
    ws.cell(row=end - 1, column=1).font = Font(bold=True)
    ws.cell(row=end - 1, column=2).font = Font(bold=True)
    row = end + 1
    key = [["Cement", t["cement_bags"], "bags"], ["Steel", t["steel_kg"], "kg"], ["Bricks", t["bricks"], "nos"],
           ["M-Sand", t["msand_m3"], "m³"], ["P-Sand", t["psand_m3"], "m³"], ["Aggregate", t["aggregate_m3"], "m³"],
           ["Labour", t["man_days"], "man-days"]]
    table(ws, row, ["Key quantity", "Qty", "Unit"], key)
    ws.freeze_panes = None

    # BOQ
    ws = wb.create_sheet("BOQ")
    rows = [[i["sl"], i["floor"], i["category"], i["description"], i["qty"], i["unit"], i["rate"],
             i["material_cost"], i["labour_cost"], i["equipment_cost"], i["amount"]] for i in r["items"]]
    end = table(ws, 1, ["Sl", "Floor", "Category", "Description", "Qty", "Unit", "Rate (Rs.)", "Material",
                        "Labour", "Equipment", "Amount (Rs.)"], rows, money_cols=(7, 8, 9, 10, 11),
                widths=[5, 15, 18, 58, 10, 8, 12, 14, 14, 12, 15])
    ws.cell(row=end, column=4, value="TOTAL (before extras)").font = Font(bold=True)
    for col in (8, 9, 10, 11):
        L = get_column_letter(col)
        c = ws.cell(row=end, column=col, value=f"=SUM({L}2:{L}{end - 1})")
        c.number_format, c.font = money, Font(bold=True)

    # Materials
    ws = wb.create_sheet("Materials")
    rows = [[m["name"], m["qty"], m["unit"], m["rate"], m["amount"]] for m in r["materials"]]
    table(ws, 1, ["Material (incl. wastage)", "Qty", "Unit", "Rate (Rs.)", "Amount (Rs.)"], rows,
          money_cols=(4, 5), widths=[50, 12, 8, 12, 16])

    # Labour
    ws = wb.create_sheet("Labour")
    rows = [[l["name"], l["qty"], l["rate"], l["amount"]] for l in r["labour"]]
    table(ws, 1, ["Trade", "Man-days", "Day wage (Rs.)", "Amount (Rs.)"], rows, money_cols=(3, 4),
          widths=[36, 12, 16, 16])

    # Equipment
    ws = wb.create_sheet("Equipment")
    rows = [[e["name"], e["qty"], e["unit"], e["rate"], e["amount"]] for e in r["equipment"]]
    table(ws, 1, ["Equipment", "Qty", "Unit", "Rate (Rs.)", "Amount (Rs.)"], rows, money_cols=(4, 5),
          widths=[40, 10, 8, 12, 16])

    # Floor-wise
    ws = wb.create_sheet("Floor-wise")
    rows = [[f["floor"], f["builtup_sqft"] or "", f["material_cost"], f["labour_cost"], f["equipment_cost"],
             f["total"]] for f in r["floors"]]
    table(ws, 1, ["Floor / Part", "Built-up (sq ft)", "Material", "Labour", "Equipment", "Total (Rs.)"], rows,
          money_cols=(3, 4, 5, 6), widths=[20, 14, 16, 16, 14, 16])

    # Rate analysis (per item breakdown)
    ws = wb.create_sheet("Rate Analysis")
    rows = []
    for i in r["items"]:
        rows.append([i["sl"], i["description"], f"{i['qty']} {i['unit']}", "", "", "", ""])
        for m in i["materials"]:
            rows.append(["", f"   Material: {m['name']}", "", m["qty"], m["unit"], m["rate"], m["amount"]])
        for l in i["labour"]:
            rows.append(["", f"   Labour: {l['name']}", "", l["days"], "days", l["rate"], l["amount"]])
        for e in i["equipment"]:
            rows.append(["", f"   Equipment: {e['name']}", "", e["qty"], e["unit"], e["rate"], e["amount"]])
    table(ws, 1, ["Sl", "Item / component", "Item qty", "Qty", "Unit", "Rate", "Amount (Rs.)"], rows,
          money_cols=(6, 7), widths=[5, 62, 14, 10, 8, 10, 14])

    # Assumptions
    ws = wb.create_sheet("Assumptions")
    ws["A1"] = "Assumptions & notes"
    ws["A1"].font = Font(bold=True, size=13, color=NAVY)
    rr = 3
    for a in r["warnings"]:
        ws.cell(row=rr, column=1, value="WARNING: " + a).font = Font(bold=True, color="B42318")
        rr += 1
    for a in r["assumptions"] + [r["note"]]:
        ws.cell(row=rr, column=1, value="• " + a)
        rr += 1
    ws.column_dimensions["A"].width = 130

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ PDF
def to_pdf(project: dict, est: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle)

    import json

    def pdf_safe(text: str) -> str:  # built-in PDF fonts lack these glyphs
        for k, v in {"²": "2", "³": "3", "≈": "~", "₹": "Rs.", "≤": "<=", "≥": ">="}.items():
            text = text.replace(k, v)
        return text

    est = json.loads(pdf_safe(json.dumps(est, ensure_ascii=False, default=str)))
    project = json.loads(pdf_safe(json.dumps(project, ensure_ascii=False, default=str)))
    r = est["result"]
    t = r["totals"]
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm,
                            topMargin=12 * mm, bottomMargin=12 * mm, title="Nexora Build Estimate")
    ss = getSampleStyleSheet()
    h1 = ParagraphStyle("h1", parent=ss["Title"], textColor=colors.HexColor("#" + NAVY), alignment=0, fontSize=18)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], textColor=colors.HexColor("#" + NAVY))
    small = ParagraphStyle("small", parent=ss["BodyText"], fontSize=8, leading=10)
    cell = ParagraphStyle("cell", parent=ss["BodyText"], fontSize=7.5, leading=9)

    def tbl(data, widths, money_from=None, bold_last=False):
        tb = Table(data, colWidths=widths, repeatRows=1)
        st = [("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#" + NAVY)),
              ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
              ("FONTSIZE", (0, 0), (-1, -1), 7.5), ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#D9DEE8")),
              ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F4F7FC")]),
              ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]
        if money_from is not None:
            st.append(("ALIGN", (money_from, 1), (-1, -1), "RIGHT"))
        if bold_last:
            st.append(("FONTNAME", (0, -1), (-1, -1), "Helvetica-Bold"))
        tb.setStyle(TableStyle(st))
        return tb

    story = [Paragraph("Nexora Build – Construction Estimate", h1), Spacer(1, 4)]
    info = [[Paragraph(f"<b>{k}</b>", small), Paragraph(pdf_safe(v), small)] for k, v in header_lines(project, est)]
    story.append(Table(info, colWidths=[40 * mm, 120 * mm], style=[("FONTSIZE", (0, 0), (-1, -1), 8)]))
    story.append(Spacer(1, 6))

    cost = [["Cost head", "Amount (Rs.)"], ["Materials", inr(t["materials"])], ["Labour", inr(t["labour"])],
            ["Equipment", inr(t["equipment"])]]
    cost += [[f"{e['name']} ({e['basis']})", inr(e["amount"])] for e in r["extras"]]
    cost += [["GRAND TOTAL", inr(t["grand_total"])]]
    keyq = [["Key quantity", "Qty"], ["Cement", f"{t['cement_bags']:,} bags"], ["Steel", f"{t['steel_kg']:,} kg"],
            ["Bricks", f"{t['bricks']:,} nos"], ["M-Sand", f"{t['msand_m3']} m3"], ["P-Sand", f"{t['psand_m3']} m3"],
            ["Aggregate", f"{t['aggregate_m3']} m3"], ["Labour", f"{t['man_days']:,.0f} man-days"]]
    floors = [["Floor / Part", "Total (Rs.)"]] + [[f["floor"], inr(f["total"])] for f in r["floors"]]
    story.append(Table([[tbl(cost, [62 * mm, 30 * mm], 1, True), tbl(keyq, [30 * mm, 32 * mm], 1),
                         tbl(floors, [34 * mm, 28 * mm], 1)]], colWidths=[96 * mm, 66 * mm, 66 * mm],
                       style=[("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story.append(Spacer(1, 6))
    if r["warnings"]:
        story.append(Paragraph("<b>Warnings</b>", h2))
        for w in r["warnings"]:
            story.append(Paragraph("• " + w, small))
    story.append(Paragraph("<b>Assumptions</b>", h2))
    for a in r["assumptions"] + [r["note"]]:
        story.append(Paragraph("• " + a, small))

    story += [PageBreak(), Paragraph("Bill of Quantities (BOQ)", h2)]
    boq = [["Sl", "Floor", "Description", "Qty", "Unit", "Rate", "Material", "Labour", "Equip.", "Amount"]]
    for i in r["items"]:
        boq.append([i["sl"], i["floor"], Paragraph(i["description"], cell), f"{i['qty']:,}", i["unit"],
                    inr(i["rate"]), inr(i["material_cost"]), inr(i["labour_cost"]), inr(i["equipment_cost"]),
                    inr(i["amount"])])
    boq.append(["", "", "TOTAL (before extras)", "", "", "", inr(t["materials"]), inr(t["labour"]),
                inr(t["equipment"]), inr(t["materials"] + t["labour"] + t["equipment"])])
    story.append(tbl(boq, [9 * mm, 24 * mm, 90 * mm, 18 * mm, 12 * mm, 18 * mm, 24 * mm, 22 * mm, 18 * mm, 26 * mm],
                     3, True))

    story += [PageBreak(), Paragraph("Materials (incl. wastage)", h2)]
    mats = [["Material", "Qty", "Unit", "Rate", "Amount"]] + [
        [m["name"], f"{m['qty']:,}", m["unit"], inr(m["rate"]), inr(m["amount"])] for m in r["materials"]]
    labs = [["Labour trade", "Man-days", "Wage", "Amount"]] + [
        [l["name"], f"{l['qty']:,}", inr(l["rate"]), inr(l["amount"])] for l in r["labour"]]
    story.append(tbl(mats, [110 * mm, 30 * mm, 18 * mm, 26 * mm, 32 * mm], 1))
    story += [Spacer(1, 8), Paragraph("Labour (man-days × TN day wage)", h2)]
    story.append(tbl(labs, [110 * mm, 30 * mm, 26 * mm, 32 * mm], 1))

    def footer(canvas, d):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.grey)
        canvas.drawString(12 * mm, 7 * mm, "Nexora Build · preliminary estimate – verify rates locally; "
                                            "structural sizes need a qualified engineer")
        canvas.drawRightString(285 * mm, 7 * mm, f"Page {d.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()
