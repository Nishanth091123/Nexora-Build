"""Excel input method for the estimate: download a template, fill it in Excel, upload it back.

Sheets:
  * Inputs      – building details (one row per item)
  * Extra Items – your own BOQ items: rate code + quantity (+ note)
  * Rates       – your rates (owner only: changed values update the rate library of that state)
"""
from __future__ import annotations

import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

# (key, label, kind, note)   kind: list | bool | int | float | choice:a/b/c
FIELDS = [
    ("floor_areas_sqft", "Built-up area of each floor (sq ft)", "list", "Ground first, comma separated, e.g. 1200, 1150"),
    ("stilt", "First level is stilt parking", "bool", "Yes / No"),
    ("floor_height_ft", "Floor height (ft)", "float", "Usually 10"),
    ("structure", "Structure", "choice:rcc_framed/load_bearing", "rcc_framed or load_bearing"),
    ("spec_level", "Specification", "choice:basic/standard/premium", "basic, standard or premium"),
    ("masonry", "Walls", "choice:brick/aac/solid_block", "brick, aac or solid_block"),
    ("bedrooms", "Bedrooms (total)", "int", ""),
    ("halls", "Halls (total)", "int", ""),
    ("kitchens", "Kitchens (total)", "int", ""),
    ("toilets", "Toilets (total)", "int", ""),
    ("plot_width_ft", "Plot width (ft)", "float", "Needed for compound wall"),
    ("plot_length_ft", "Plot length (ft)", "float", ""),
    ("staircase", "Staircase", "bool", "Yes / No"),
    ("compound_wall", "Compound wall", "bool", "Yes / No"),
    ("gate", "Gate", "bool", "Yes / No"),
    ("sump_litres", "Sump (litres)", "int", "0 = none"),
    ("septic_tank", "Septic tank", "bool", "Yes / No"),
    ("oh_tank_litres", "Overhead tank (litres)", "int", "Blank = 1000 per kitchen"),
    ("rwh", "Rain water harvesting", "bool", "Yes / No"),
    ("parking_area_sqft", "Paved parking area (sq ft)", "float", ""),
    ("approval_fees", "Plan approval & fees (Rs)", "float", ""),
    ("contractor_margin_pct", "Contractor margin (%)", "float", "0 = owner-managed"),
    ("gst_pct", "GST (%)", "float", ""),
]
HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="1D4ED8")
NOTE = Font(italic=True, color="64748B")


class ExcelInputError(ValueError):
    pass


def _fmt(kind: str, v):
    if v is None:
        return ""
    if kind == "list":
        return ", ".join(f"{float(x):g}" for x in (v or []))
    if kind == "bool":
        return "Yes" if v else "No"
    return v


def template(inputs: dict, rates: list[dict], state: str, project: str, is_owner: bool) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Inputs"
    ws.append([f"Nexora Build – estimate inputs · {project}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Fill column B, save, then upload this file in Estimate → Input Methods → Excel. Do not change column A."])
    ws["A2"].font = NOTE
    ws.append(["Key", "Value", "Item", "Note"])
    for c in ws[3]:
        c.font, c.fill = HEAD, FILL
    for key, label, kind, note in FIELDS:
        ws.append([key, _fmt(kind, inputs.get(key)), label, note])
    yes_no = DataValidation(type="list", formula1='"Yes,No"', allow_blank=True)
    ws.add_data_validation(yes_no)
    for i, (key, _, kind, _) in enumerate(FIELDS, start=4):
        if kind == "bool":
            yes_no.add(f"B{i}")
        elif kind.startswith("choice:"):
            dv = DataValidation(type="list", formula1='"' + ",".join(kind[7:].split("/")) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"B{i}")
    for col, w in zip("ABCD", (24, 22, 38, 40)):
        ws.column_dimensions[col].width = w

    ex = wb.create_sheet("Extra Items")
    ex.append(["Rate code", "Quantity", "Note", "", "Rate codes you can use (see the Rates sheet)"])
    for c in ex[1][:3]:
        c.font, c.fill = HEAD, FILL
    for x in inputs.get("extra_items") or []:
        ex.append([x.get("code"), x.get("qty"), x.get("note", "")])
    for col, w in zip("ABCE", (18, 12, 40, 50)):
        ex.column_dimensions[col].width = w

    rs = wb.create_sheet("Rates")
    rs.append([f"Rates – {state}" + ("" if is_owner else " (view only – only the owner can change rates)")])
    rs["A1"].font = Font(bold=True)
    rs.append(["Code", "Item", "Unit", "Basic", "Standard", "Premium", "Category"])
    for c in rs[2]:
        c.font, c.fill = HEAD, FILL
    for r in rates:
        rs.append([r["code"], r["name"], r["unit"], r["basic"], r["standard"], r["premium"], r.get("category", "")])
    for col, w in zip("ABCDEFG", (18, 46, 10, 12, 12, 12, 18)):
        rs.column_dimensions[col].width = w
    for row in rs.iter_rows(min_row=3, min_col=4, max_col=6):
        for c in row:
            c.number_format = "#,##0.00"
            c.alignment = Alignment(horizontal="right")
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _bool(v) -> bool:
    return str(v).strip().lower() in ("yes", "y", "true", "1", "aam", "ஆம்")


def _num(v, key):
    if v in (None, ""):
        return None
    try:
        return float(str(v).replace(",", "").strip())
    except ValueError:
        raise ExcelInputError(f"'{key}' must be a number (found '{v}')")


def parse(data: bytes, known_codes: set[str]) -> dict:
    """Returns {"inputs": {...}, "extra_items": [...], "rates": {code: {basic, standard, premium}}, "warnings": [...]}."""
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception:
        raise ExcelInputError("This is not a valid .xlsx file – download the template and fill that.")
    if "Inputs" not in wb.sheetnames:
        raise ExcelInputError("Sheet 'Inputs' not found – please use the Nexora template.")
    kinds = {k: kind for k, _, kind, _ in FIELDS}
    inputs, warnings = {}, []
    for row in wb["Inputs"].iter_rows(min_row=4, values_only=True):
        if not row or row[0] not in kinds:
            continue
        key, val = row[0], row[1] if len(row) > 1 else None
        kind = kinds[key]
        if val in (None, ""):
            continue
        if kind == "list":
            parts = [p for p in str(val).replace(";", ",").split(",") if p.strip()]
            areas = [_num(p, key) for p in parts]
            if not areas or any(a is None or a <= 0 for a in areas):
                raise ExcelInputError("Built-up area of each floor must be numbers like: 1200, 1150")
            inputs[key] = areas
        elif kind == "bool":
            inputs[key] = _bool(val)
        elif kind == "int":
            inputs[key] = int(_num(val, key))
        elif kind == "float":
            inputs[key] = _num(val, key)
        else:
            opts = kind[7:].split("/")
            v = str(val).strip().lower().replace(" ", "_")
            if v not in opts:
                raise ExcelInputError(f"'{key}' must be one of: {', '.join(opts)} (found '{val}')")
            inputs[key] = v
    extra = []
    if "Extra Items" in wb.sheetnames:
        for row in wb["Extra Items"].iter_rows(min_row=2, values_only=True):
            if not row or not row[0]:
                continue
            code = str(row[0]).strip()
            qty = _num(row[1] if len(row) > 1 else None, f"quantity of {code}")
            if code not in known_codes:
                warnings.append(f"Extra item '{code}' is not in the rate library – skipped.")
                continue
            if not qty or qty <= 0:
                warnings.append(f"Extra item '{code}' has no quantity – skipped.")
                continue
            extra.append({"code": code, "qty": qty, "note": str(row[2] or "") if len(row) > 2 else ""})
    rates = {}
    if "Rates" in wb.sheetnames:
        for row in wb["Rates"].iter_rows(min_row=3, values_only=True):
            if not row or not row[0] or str(row[0]).strip() not in known_codes:
                continue
            vals = {}
            for i, lvl in ((3, "basic"), (4, "standard"), (5, "premium")):
                v = _num(row[i] if len(row) > i else None, f"{row[0]} {lvl} rate")
                if v is not None:
                    if v < 0:
                        raise ExcelInputError(f"Rate of {row[0]} cannot be negative")
                    vals[lvl] = v
            rates[str(row[0]).strip()] = vals
    return {"inputs": inputs, "extra_items": extra, "rates": rates, "warnings": warnings}
