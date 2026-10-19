"""Site reports: Daily Progress Report (PDF) and period summary (Excel)."""
from __future__ import annotations

import io
from datetime import date, timedelta

from sqlalchemy.orm import Session

from ..design.exports import _pdf_safe
from ..models import (Attendance, MaterialEntry, Project, QualityCheck, SafetyRecord, SiteActivity, SitePhoto, SiteReport,
                      StoredFile, Worker)
from ..storage import backend_for
from . import logic


def _stages(db, pid):
    from ..routers.pm import as_dicts, stages_of
    rows = stages_of(db, pid)
    return rows, as_dicts(rows)


def _trades(db, p):
    from ..routers.site import trade_rates
    from ..estimate.rates_seed import DEFAULT_STATE
    return trade_rates(db, p.state or DEFAULT_STATE)


def dpr_pdf(db: Session, p: Project, d: date) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    ss = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=ss["Normal"], fontSize=9, leading=12)
    small = ParagraphStyle("s", parent=body, fontSize=8, leading=10, textColor=colors.HexColor("#475569"))
    h1 = ParagraphStyle("h1", parent=ss["Heading1"], fontSize=15, spaceAfter=2)
    h2 = ParagraphStyle("h2", parent=ss["Heading2"], fontSize=11, spaceBefore=8, spaceAfter=3)
    P = lambda t, st=body: Paragraph(_pdf_safe(str(t)).replace("\n", "<br/>"), st)

    def table(rows, widths):
        t = Table([[P(c) if not hasattr(c, "wrap") else c for c in r] for r in rows], colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef2ff")),
                               ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#cbd5e1")), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        return t

    rep = db.query(SiteReport).filter_by(project_id=p.id, report_date=d).first()
    srows, sd = _stages(db, p.id)
    story = [P(f"Daily Progress Report – {p.name}", h1),
             P(f"{p.code} · {p.location or ''} · Date: {d:%d %b %Y} ({d:%A})", small), Spacer(1, 4)]
    if sd:
        story.append(P(f"<b>Overall progress {logic.overall(sd):g}%</b> (planned {logic.planned(sd, d):g}%) – "
                       f"{'ahead' if logic.schedule_variance_days(sd, d) >= 0 else 'behind'} by "
                       f"{abs(logic.schedule_variance_days(sd, d))} day(s)."))
    if rep:
        story += [P(f"Weather: {rep.weather or '—'}", small), P("Work done today", h2), P(rep.work_done or "—"),
                  P("Plan for tomorrow", h2), P(rep.next_plan or "—")]
        if rep.remarks:
            story += [P("Remarks", h2), P(rep.remarks)]
    acts = (db.query(SiteActivity).filter(SiteActivity.project_id == p.id, SiteActivity.work_date == d,
                                          SiteActivity.is_hidden.is_(False)).all())
    story.append(P("Activities", h2))
    story.append(table([["Activity", "Location", "Status", "Progress"]] +
                       [[a.title, a.location or "—", a.status.replace("_", " "), f"{a.progress:g}%"] for a in acts]
                       or [["Activity", "Location", "Status", "Progress"], ["—", "", "", ""]],
                       [75 * mm, 40 * mm, 30 * mm, 25 * mm]) if acts else P("No activities recorded.", small))
    trades = _trades(db, p)
    att = (db.query(Attendance, Worker).join(Worker, Worker.id == Attendance.worker_id)
           .filter(Attendance.project_id == p.id, Attendance.work_date == d).all())
    story.append(P("Manpower", h2))
    if att:
        by: dict = {}
        for a, w in att:
            if a.status == "absent":
                continue
            by.setdefault(trades.get(w.trade, {}).get("name", w.trade), []).append(w.name + (" (½)" if a.status == "half" else ""))
        story.append(table([["Trade", "Count", "Names"]] + [[k, str(len(v)), ", ".join(v)] for k, v in by.items()],
                           [50 * mm, 18 * mm, 102 * mm]))
    else:
        story.append(P("Attendance not marked.", small))
    mats = db.query(MaterialEntry).filter(MaterialEntry.project_id == p.id, MaterialEntry.entry_date == d,
                                          MaterialEntry.is_hidden.is_(False)).all()
    story.append(P("Materials", h2))
    story.append(table([["Material", "Received / used", "Qty", "Supplier / bill"]] +
                       [[m.name, m.kind, f"{m.qty:g} {m.unit}", f"{m.supplier} {m.bill_no}".strip() or "—"] for m in mats],
                       [60 * mm, 30 * mm, 30 * mm, 50 * mm]) if mats else P("No material entries.", small))
    qcs = db.query(QualityCheck).filter(QualityCheck.project_id == p.id, QualityCheck.checked_on == d).all()
    if qcs:
        story += [P("Quality checks", h2), table([["Check", "Result", "Remarks"]] + [[q.item, q.status, q.remarks or "—"] for q in qcs],
                                                 [80 * mm, 25 * mm, 65 * mm])]
    saf = db.query(SafetyRecord).filter(SafetyRecord.project_id == p.id, SafetyRecord.record_date == d,
                                        SafetyRecord.is_hidden.is_(False)).all()
    if saf:
        story += [P("Safety & issues", h2), table([["Type", "Title", "Severity", "Status"]] +
                                                  [[s.kind.replace("_", " "), s.title, s.severity, s.status] for s in saf],
                                                  [30 * mm, 90 * mm, 25 * mm, 25 * mm])]
    photos = (db.query(SitePhoto).filter(SitePhoto.project_id == p.id, SitePhoto.taken_on == d, SitePhoto.is_hidden.is_(False))
              .limit(6).all())
    if photos:
        story.append(P("Site photos", h2))
        cells = []
        for ph in photos:
            f = db.get(StoredFile, ph.file_id)
            try:
                img = Image(io.BytesIO(backend_for(f.storage_backend).read(f.storage_key)), width=55 * mm, height=40 * mm, kind="proportional")
                cells.append([img, P(ph.caption, small)])
            except Exception:
                cells.append([P("(photo)", small), P(ph.caption, small)])
        grid = [cells[i:i + 3] for i in range(0, len(cells), 3)]
        flat = []
        for row in grid:
            flat.append([c[0] for c in row] + [""] * (3 - len(row)))
            flat.append([c[1] for c in row] + [""] * (3 - len(row)))
        t = Table(flat, colWidths=[58 * mm] * 3)
        t.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
        story.append(t)
    story += [Spacer(1, 10), P("Prepared with Nexora Build. Signature: ____________________", small)]
    buf = io.BytesIO()
    SimpleDocTemplate(buf, pagesize=A4, leftMargin=14 * mm, rightMargin=14 * mm, topMargin=12 * mm, bottomMargin=12 * mm,
                      title=f"DPR {d}").build(story)
    return buf.getvalue()


def period_excel(db: Session, p: Project, start: date, end: date) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor="1D4ED8")

    def sheet(title, headers, rows, first=False):
        ws = wb.active if first else wb.create_sheet()
        ws.title = title
        ws.append([f"{p.name} – {title} ({start:%d %b %Y} to {end:%d %b %Y})"])
        ws["A1"].font = Font(bold=True, size=12)
        ws.append(headers)
        for c in ws[2]:
            c.font, c.fill = head, fill
        for r in rows:
            ws.append(list(r))
        for i, h in enumerate(headers, 1):
            width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows if i - 1 < len(r)] + [8])
            ws.column_dimensions[get_column_letter(i)].width = min(60, width + 2)
        ws.freeze_panes = "A3"
        return ws

    srows, sd = _stages(db, p.id)
    sheet("Schedule", ["Stage", "Group", "Start", "End", "Progress %", "Weight", "Milestone"],
          [[s.name, s.stage_group, s.start_date, s.end_date, s.progress, s.weight, "Yes" if s.is_milestone else ""] for s in srows],
          first=True)
    acts = (db.query(SiteActivity).filter(SiteActivity.project_id == p.id, SiteActivity.is_hidden.is_(False),
                                          SiteActivity.work_date >= start, SiteActivity.work_date <= end)
            .order_by(SiteActivity.work_date, SiteActivity.id).all())
    sheet("Activities", ["Date", "Activity", "Location", "Status", "Progress %", "Notes"],
          [[a.work_date, a.title, a.location, a.status, a.progress, a.notes] for a in acts])
    reps = (db.query(SiteReport).filter(SiteReport.project_id == p.id, SiteReport.report_date >= start,
                                        SiteReport.report_date <= end).order_by(SiteReport.report_date).all())
    sheet("Daily Reports", ["Date", "Weather", "Work done", "Next plan", "Remarks"],
          [[r.report_date, r.weather, r.work_done, r.next_plan, r.remarks] for r in reps])
    mats = (db.query(MaterialEntry).filter(MaterialEntry.project_id == p.id, MaterialEntry.is_hidden.is_(False),
                                           MaterialEntry.entry_date >= start, MaterialEntry.entry_date <= end)
            .order_by(MaterialEntry.entry_date, MaterialEntry.id).all())
    sheet("Materials", ["Date", "Material", "Received / used", "Qty", "Unit", "Rate", "Amount", "Supplier", "Bill no"],
          [[m.entry_date, m.name, m.kind, m.qty, m.unit, m.rate or "", round(m.qty * m.rate, 2) if m.rate else "", m.supplier, m.bill_no]
           for m in mats])
    # attendance matrix
    trades = _trades(db, p)
    workers = db.query(Worker).filter(Worker.project_id == p.id, Worker.is_hidden.is_(False)).order_by(Worker.trade, Worker.name).all()
    days = [start + timedelta(days=i) for i in range((end - start).days + 1)]
    marks = {(a.worker_id, a.work_date): a for a in db.query(Attendance).filter(
        Attendance.project_id == p.id, Attendance.work_date >= start, Attendance.work_date <= end).all()}
    rows = []
    for w in workers:
        wage = w.day_wage if w.day_wage is not None else trades.get(w.trade, {}).get("wage", 0) or 0
        cells, tot_days, pay = [], 0.0, 0.0
        for d in days:
            a = marks.get((w.id, d))
            code = {"present": "P", "half": "H", "absent": "A"}.get(a.status, "") if a else ""
            f = {"P": 1.0, "H": 0.5}.get(code, 0.0)
            tot_days += f
            wg = a.wage if a and a.wage is not None else wage
            pay += f * wg + ((a.ot_hours or 0) * wg / 8 if a else 0)
            cells.append(code)
        if tot_days or w.is_active:
            rows.append([w.name, trades.get(w.trade, {}).get("name", w.trade), wage] + cells + [tot_days, round(pay)])
    ws = sheet("Attendance", ["Worker", "Trade", "Day wage"] + [d.strftime("%d") for d in days] + ["Days", "Wages ₹"], rows)
    for col in range(4, 4 + len(days)):
        ws.column_dimensions[get_column_letter(col)].width = 4
        for r in ws.iter_rows(min_row=3, min_col=col, max_col=col):
            for c in r:
                c.alignment = Alignment(horizontal="center")
    qcs = db.query(QualityCheck).filter(QualityCheck.project_id == p.id, QualityCheck.is_hidden.is_(False)).order_by(QualityCheck.sort).all()
    sheet("Quality", ["Stage", "Check", "Standard", "Status", "Checked on", "Remarks"],
          [[q.stage, q.item, q.standard, q.status, q.checked_on or "", q.remarks] for q in qcs])
    saf = (db.query(SafetyRecord).filter(SafetyRecord.project_id == p.id, SafetyRecord.is_hidden.is_(False),
                                         SafetyRecord.record_date >= start, SafetyRecord.record_date <= end)
           .order_by(SafetyRecord.record_date).all())
    sheet("Safety & Issues", ["Date", "Type", "Title", "Severity", "Status", "Delay days", "Action", "Closed on"],
          [[s.record_date, s.kind, s.title, s.severity, s.status, s.delay_days, s.action, s.closed_on or ""] for s in saf])
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=3):
            for c in row:
                if isinstance(c.value, date):
                    c.number_format = "DD-MMM-YYYY"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
