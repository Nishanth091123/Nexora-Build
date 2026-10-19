"""BBS exports: Excel (schedule + steel summary) and PDF (with bar-shape sketches and member drawings)."""
from __future__ import annotations

import io
from datetime import datetime


def to_excel(project: dict, title: str, res: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "BBS"
    ws.append([f"Bar Bending Schedule – {title}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"{project.get('name', '')} · {project.get('code', '')} · {datetime.now():%d %b %Y}"])
    ws.append([])
    head = ["Member", "Bar mark", "Description", "Shape", "Shape sizes (mm)", "Dia (mm)", "Nos / member", "Members",
            "Total nos", "Cutting length (mm)", "Total length (m)", "Weight (kg)"]
    ws.append(head)
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1D4ED8")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for r in res["rows"]:
        dims = ", ".join(f"{k}={v}" for k, v in r["dims"].items())
        ws.append([r["member"], r["mark"], r["desc"], r["shape"], dims, r["dia"], r["per_member"], r["members"], r["nos"],
                   r["cut_mm"], r["total_m"], r["kg"]])
    for col, w in zip("ABCDEFGHIJKL", (9, 12, 40, 9, 22, 8, 11, 9, 9, 13, 13, 11)):
        ws.column_dimensions[col].width = w
    s = wb.create_sheet("Steel summary")
    s.append(["Diameter (mm)", "kg / m", "Total length (m)", "Weight (kg)", f"With {res['wastage_pct']:g}% wastage (kg)",
              "12 m bars to order"])
    for c in s[1]:
        c.font = Font(bold=True)
    for x in res["summary"]:
        s.append([x["dia"], x["kg_per_m"], x["total_m"], x["kg"], x["kg_with_wastage"], x["bars_12m"]])
    s.append(["Total", "", "", round(sum(x["kg"] for x in res["summary"]), 1), res["total_kg"], ""])
    s.append([])
    for a in res["assumptions"]:
        s.append([a])
    for col, w in zip("ABCDEF", (14, 9, 16, 13, 22, 17)):
        s.column_dimensions[col].width = w
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _shape(r: dict, w=110, h=34):
    """Small bar-shape sketch for one schedule row."""
    from reportlab.graphics.shapes import Drawing, Line, PolyLine, String
    from reportlab.lib import colors
    d = Drawing(w, h)
    blue = colors.HexColor("#1d4ed8")
    s, dims = r["shape"], r["dims"]
    if s == "straight":
        d.add(Line(8, h / 2, w - 8, h / 2, strokeColor=blue, strokeWidth=2))
        d.add(String(w / 2, h / 2 + 4, str(dims.get("a", "")), fontSize=6, textAnchor="middle"))
    elif s == "L":
        d.add(PolyLine([8, 6, 8, h - 6, w - 8, h - 6], strokeColor=blue, strokeWidth=2))
        d.add(String(12, 10, str(dims.get("a", "")), fontSize=6))
        d.add(String(w / 2, h - 14, str(dims.get("b", "")), fontSize=6, textAnchor="middle"))
    elif s == "U":
        d.add(PolyLine([8, h - 6, 8, 8, w - 8, 8, w - 8, h - 6], strokeColor=blue, strokeWidth=2))
        d.add(String(w / 2, 12, str(dims.get("b", "")), fontSize=6, textAnchor="middle"))
        d.add(String(11, h - 12, str(dims.get("a", "")), fontSize=6))
    elif s == "stirrup":
        a, b = dims.get("a", 1) or 1, dims.get("b", 1) or 1
        k = min((w - 30) / a, (h - 8) / b)
        x0, y0, ww, hh = 10, 4, a * k, b * k
        d.add(PolyLine([x0, y0, x0 + ww, y0, x0 + ww, y0 + hh, x0, y0 + hh, x0, y0], strokeColor=blue, strokeWidth=2))
        d.add(Line(x0, y0 + hh, x0 + 6, y0 + hh - 6, strokeColor=blue, strokeWidth=2))
        d.add(Line(x0 + 3, y0 + hh, x0 + 9, y0 + hh - 6, strokeColor=blue, strokeWidth=2))
        d.add(String(x0 + ww + 4, y0 + hh / 2, f"{dims.get('a')}x{dims.get('b')}", fontSize=6))
    return d


def _member(dr: dict, w=250, h=200):
    """Simple member drawing: column elevation (lap zone + tie zones), beam elevation, footing / slab plan."""
    from reportlab.graphics.shapes import Drawing, Line, Rect, String
    from reportlab.lib import colors
    d = Drawing(w, h)
    blue, red, grey, amber = (colors.HexColor("#1d4ed8"), colors.HexColor("#dc2626"), colors.HexColor("#94a3b8"),
                              colors.HexColor("#fde68a"))
    t = dr["type"]
    if t == "column":
        floors, fh = dr["floors"], dr["h"]
        total = dr["footing_depth"] + floors * fh
        k = (h - 30) / total
        x, cw = 40, 40
        y = 15
        d.add(Rect(x - 30, y, cw + 60, dr["footing_depth"] * k, strokeColor=grey, fillColor=None))
        for i in range(floors):
            y0 = y + (dr["footing_depth"] + i * fh) * k
            d.add(Line(x - 10, y0, x + cw + 10, y0, strokeColor=grey))
            if i < floors - 1:  # lap in the middle half of the floor
                ly = y0 + fh * k * 0.25
                d.add(Rect(x, ly, cw, fh * k * 0.5, strokeColor=None, fillColor=amber))
                d.add(String(x + cw + 14, ly + fh * k * 0.25, f"lap {dr['lap_mm']} mm here", fontSize=6, fillColor=red))
            # close ties at both ends
            for frac in (0.0, 1.0):
                zy = y0 + (0 if frac == 0 else fh * k - dr["l0"] * k)
                for j in range(4):
                    d.add(Line(x, zy + j * dr["l0"] * k / 4, x + cw, zy + j * dr["l0"] * k / 4, strokeColor=blue, strokeWidth=0.6))
        d.add(Line(x + 6, y + 10, x + 6, h - 12, strokeColor=blue, strokeWidth=1.5))
        d.add(Line(x + cw - 6, y + 10, x + cw - 6, h - 12, strokeColor=blue, strokeWidth=1.5))
        d.add(Line(x + 6, y + 10, x - 14, y + 10, strokeColor=blue, strokeWidth=1.5))
        d.add(Line(x + cw - 6, y + 10, x + cw + 14, y + 10, strokeColor=blue, strokeWidth=1.5))
        d.add(String(x + cw + 14, h - 20, f"{dr['b']:g}x{dr['d']:g}, ties {dr['tie_dia']:g} mm", fontSize=7))
        d.add(String(x + cw + 14, h - 30, f"@ {dr['end_spacing']:g} mm for {dr['l0']} mm at ends", fontSize=6))
        d.add(String(x + cw + 14, h - 40, f"@ {dr['spacing']:g} mm in the middle", fontSize=6))
        d.add(String(x + cw + 14, y + 4, f"L-bend {dr['bottom_bend']:g} mm in footing", fontSize=6))
    elif t == "beam":
        k = (w - 30) / (dr["span"] + 2 * dr["support"])
        x0, y0, bh = 15, h / 2 - 20, 40
        d.add(Rect(x0, y0 - 10, dr["support"] * k, bh + 20, fillColor=colors.HexColor("#e2e8f0"), strokeColor=grey))
        d.add(Rect(x0 + (dr["support"] + dr["span"]) * k, y0 - 10, dr["support"] * k, bh + 20,
                   fillColor=colors.HexColor("#e2e8f0"), strokeColor=grey))
        d.add(Rect(x0, y0, (dr["span"] + 2 * dr["support"]) * k, bh, strokeColor=grey, fillColor=None))
        d.add(Line(x0 + 4, y0 + 5, x0 + (dr["span"] + 2 * dr["support"]) * k - 4, y0 + 5, strokeColor=blue, strokeWidth=1.5))
        d.add(Line(x0 + 4, y0 + bh - 5, x0 + (dr["span"] + 2 * dr["support"]) * k - 4, y0 + bh - 5, strokeColor=blue, strokeWidth=1.5))
        s0 = x0 + dr["support"] * k
        for i in range(30):
            xx = s0 + i * (dr["span"] * k) / 29
            d.add(Line(xx, y0 + 3, xx, y0 + bh - 3, strokeColor=blue, strokeWidth=0.4))
        d.add(String(w / 2, y0 - 22, f"span {dr['span']:g} mm; stirrups @ {dr['end_spacing']} near supports "
                     f"({dr['zone']} mm), @ {dr['spacing']:g} middle", fontSize=6, textAnchor="middle"))
    else:
        L = dr.get("length") or dr.get("ly")
        B = dr.get("width") or dr.get("lx")
        k = min((w - 20) / L, (h - 30) / B)
        x0, y0 = 10, 20
        d.add(Rect(x0, y0, L * k, B * k, strokeColor=grey, fillColor=None))
        n1 = min(int(dr.get("nl") or dr.get("nd") or 8), 25)
        n2 = min(int(dr.get("nb") or dr.get("nm") or 8), 25)
        for i in range(n1):
            yy = y0 + 3 + i * (B * k - 6) / max(1, n1 - 1)
            d.add(Line(x0 + 3, yy, x0 + L * k - 3, yy, strokeColor=blue, strokeWidth=0.6))
        for i in range(n2):
            xx = x0 + 3 + i * (L * k - 6) / max(1, n2 - 1)
            d.add(Line(xx, y0 + 3, xx, y0 + B * k - 3, strokeColor=red, strokeWidth=0.6))
        d.add(String(x0, 6, f"{L:g} x {B:g} mm", fontSize=7))
    return d


def to_pdf(project: dict, title: str, res: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm,
                            bottomMargin=12 * mm, title=f"BBS – {title}")
    st = getSampleStyleSheet()
    small = st["BodyText"].clone("s", fontSize=7.5, leading=9)
    el = [Paragraph(f"<b>Bar Bending Schedule – {title}</b>", st["Title"]),
          Paragraph(f"{project.get('name', '')} · {project.get('code', '')} · {datetime.now():%d %b %Y}", small), Spacer(1, 6)]
    head = ["Member", "Mark", "Description", "Shape", "Dia", "Nos/mem", "Mem", "Total nos", "Cut length (mm)",
            "Total (m)", "Weight (kg)"]
    data = [head]
    for r in res["rows"]:
        data.append([r["member"], r["mark"], Paragraph(r["desc"], small), _shape(r), f"{r['dia']:g}", r["per_member"],
                     r["members"], r["nos"], f"{r['cut_mm']:,}", f"{r['total_m']:,.1f}", f"{r['kg']:,.1f}"])
    t = Table(data, colWidths=[16 * mm, 20 * mm, 62 * mm, 42 * mm, 11 * mm, 15 * mm, 11 * mm, 16 * mm, 24 * mm, 20 * mm, 22 * mm],
              repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1d4ed8")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                           ("FONTSIZE", (0, 0), (-1, -1), 7.5), ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#cbd5e1")),
                           ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
    el += [t, Spacer(1, 8), Paragraph("<b>Steel summary</b>", st["Heading3"])]
    sdata = [["Dia (mm)", "kg/m", "Total length (m)", "Weight (kg)", f"+{res['wastage_pct']:g}% wastage (kg)", "12 m bars"]]
    for x in res["summary"]:
        sdata.append([f"{x['dia']:g}", x["kg_per_m"], f"{x['total_m']:,.1f}", f"{x['kg']:,.1f}", f"{x['kg_with_wastage']:,.1f}", x["bars_12m"]])
    sdata.append(["Total", "", "", "", f"{res['total_kg']:,.1f}", ""])
    s = Table(sdata)
    s.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#cbd5e1")), ("FONTSIZE", (0, 0), (-1, -1), 8),
                           ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e2e8f0"))]))
    el += [s, Spacer(1, 8), Paragraph("<b>Member drawings</b> (laps, tie / stirrup zones, mesh)", st["Heading3"])]
    cells = [[Paragraph(f"<b>{dr['mark']}</b> – {dr['type']}", small), _member(dr)] for dr in res["drawings"]]
    grid, row = [], []
    for c in cells:
        row.append(c)
        if len(row) == 3:
            grid.append(row)
            row = []
    if row:
        grid.append(row + [""] * (3 - len(row)))
    if grid:
        g = Table([[Table([[c[0]], [c[1]]]) if c else "" for c in r] for r in grid])
        g.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
        el.append(g)
    el.append(Spacer(1, 6))
    for a in res["assumptions"]:
        el.append(Paragraph("• " + a, small))
    doc.build(el)
    return buf.getvalue()
